import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = openpyxl.Workbook()

# Sheet 1: Module 7.1 & 7.2 (Phong hop & Sinh nhat)
ws = wb.active
ws.title = "CSDL_Nhom7"
ws.views.sheetView[0].showGridLines = True

# Colors & Styles
FONT_NOTE = Font(name="Calibri", size=10, italic=True, color="C00000") # Red note
FONT_SECTION = Font(name="Calibri", size=11, bold=True, color="000000")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BODY = Font(name="Calibri", size=11, color="000000")
FONT_CONSTRAINT = Font(name="Calibri", size=10, italic=True, color="333333")

FILL_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid") # Dark Navy Blue

THIN_BORDER = Border(
    left=Side(style="thin", color="D9D9D9"),
    right=Side(style="thin", color="D9D9D9"),
    top=Side(style="thin", color="D9D9D9"),
    bottom=Side(style="thin", color="D9D9D9")
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")

# Row 1: Note
ws.cell(row=1, column=1, value="Nguồn: schema-toan-du-an.md (Nhóm 7 - Phòng họp & Sinh nhật CBNV). Phụ trách: Trần Đình Tình. Đồng bộ chuẩn khóa ngoại với accounts & organizations.").font = FONT_NOTE

current_row = 3

tables = [
    {
        "title": "Bảng: Phòng họp (danh mục)",
        "constraint": "(*) Khóa chính: id. Ràng buộc: tên phòng là duy nhất (unique), sức chứa phải > 0.",
        "columns": [
            ("ID", "id", "Number (auto)", ""),
            ("Tên phòng", "name", "Text (100)", "X"),
            ("Vị trí", "location", "Text (200)", "X"),
            ("Sức chứa", "capacity", "Number", "X"),
            ("Tiện ích / Thiết bị", "amenities", "Text", ""),
            ("Trạng thái", "status", "Text (20)", "X (Mặc định: AVAILABLE)"),
            ("Thời điểm tạo", "created_at", "DateTime", "X (auto_now_add)"),
        ]
    },
    {
        "title": "Bảng: Đặt phòng họp (lịch đặt)",
        "constraint": "(*) Ràng buộc nghiệp vụ: 1 phòng + 1 khung giờ chỉ 1 lượt đặt hợp lệ (chặn trùng lịch start_time - end_time ở DB/clean logic). Khi hủy yêu cầu xác thực email người đặt.",
        "columns": [
            ("ID", "id", "Number (auto)", ""),
            ("Phòng họp", "room_id", "FK -> Phòng họp", "X"),
            ("Người đặt", "organizer_id", "FK -> CBNV", "X (hoặc NULL nếu khách vãng lai)"),
            ("Đơn vị tổ chức", "department_id", "FK -> Phòng ban", "X (hoặc NULL)"),
            ("Tiêu đề cuộc họp", "title", "Text (255)", "X"),
            ("Ngày họp", "date", "Date", "X"),
            ("Giờ bắt đầu", "start_time", "Time", "X"),
            ("Giờ kết thúc", "end_time", "Time", "X (sau giờ bắt đầu)"),
            ("Số người tham dự", "attendees_count", "Number", "X (Mặc định: 1)"),
            ("Mục đích / Ghi chú", "purpose", "Text", ""),
            ("Trạng thái", "status", "Text (20)", "X (CONFIRMED / CANCELLED / COMPLETED)"),
            ("Thời điểm tạo", "created_at", "DateTime", "X (auto_now_add)"),
        ]
    },
    {
        "title": "Bảng: Cán bộ nhân viên (hồ sơ nhân sự)",
        "constraint": "(*) Bảng dữ liệu dùng chung. Khóa chính: id. employee_code & email là duy nhất (unique). Cung cấp ngày sinh nhật cho phân hệ sinh nhật.",
        "columns": [
            ("ID", "id", "Number (auto)", ""),
            ("Mã cán bộ", "employee_code", "Text (20)", "X (Unique)"),
            ("Họ và tên", "full_name", "Text (150)", "X"),
            ("Ngày sinh", "birth_date", "Date", "X"),
            ("Email", "email", "Text (254)", "X (Unique)"),
            ("Số điện thoại", "phone", "Text (20)", ""),
            ("Phòng ban", "department_id", "FK -> Phòng ban", ""),
            ("Chức vụ", "position", "Text (100)", "X"),
            ("Đang làm việc", "is_active", "Boolean", "X (Mặc định: True)"),
            ("Thời điểm tạo", "created_at", "DateTime", "X (auto_now_add)"),
        ]
    },
    {
        "title": "Bảng: Lời chúc sinh nhật",
        "constraint": "(*) Khóa chính: id. Lưu trữ thiệp mừng và lời chúc gửi đến cán bộ nhân viên.",
        "columns": [
            ("ID", "id", "Number (auto)", ""),
            ("CBNV nhận chúc", "employee_id", "FK -> CBNV", "X"),
            ("Người / Đơn vị gửi", "sender_name", "Text (150)", "X"),
            ("Nội dung lời chúc", "message", "Text", "X"),
            ("Mẫu thiệp", "card_template", "Text (50)", "X (Mặc định: GOLDEN)"),
            ("Thời điểm gửi", "created_at", "DateTime", "X (auto_now_add)"),
        ]
    },
    {
        "title": "Bảng: Nhật ký gửi email tự động (lịch sử gửi)",
        "constraint": "(*) Ràng buộc duy nhất: (employee_id, sent_date) -> unique_together. Đảm bảo 1 nhân viên chỉ nhận tối đa 1 email tự động mỗi ngày sinh nhật.",
        "columns": [
            ("ID", "id", "Number (auto)", ""),
            ("CBNV nhận mail", "employee_id", "FK -> CBNV", "X"),
            ("Ngày sinh nhật", "sent_date", "Date", "X"),
            ("Thời điểm gửi", "sent_at", "DateTime", "X"),
            ("Email người nhận", "recipient_email", "Text (254)", "X"),
            ("Tiêu đề email", "subject", "Text (255)", "X"),
            ("Mẫu thiệp đính kèm", "card_template", "Text (50)", "X (GOLDEN)"),
            ("Trạng thái", "status", "Text (20)", "X (SUCCESS / FAILED)"),
            ("Thông báo lỗi (nếu có)", "error_message", "Text", ""),
        ]
    },
    {
        "title": "Bảng: Đơn vị / Phòng ban (cơ cấu tổ chức)",
        "constraint": "(*) Bảng danh mục đơn vị dùng chung. Mã phòng ban code là duy nhất (unique).",
        "columns": [
            ("ID", "id", "Number (auto)", ""),
            ("Mã phòng ban", "code", "Text (50)", "X (Unique)"),
            ("Tên phòng ban", "name", "Text (255)", "X"),
            ("Cấp đơn vị", "level", "Number", "X (Mặc định: 1)"),
            ("Thứ tự sắp xếp", "sort_order", "Number", "X (Mặc định: 0)"),
            ("Thời điểm tạo", "created_at", "DateTime", "X (auto_now_add)"),
        ]
    }
]

headers = ["Tên trường", "Field", "Kiểu dữ liệu", "Bắt buộc"]

for table in tables:
    # Title Row
    cell_title = ws.cell(row=current_row, column=1, value=table["title"])
    cell_title.font = FONT_SECTION
    current_row += 1
    
    # Header Row
    for col_idx, h_text in enumerate(headers, start=1):
        cell = ws.cell(row=current_row, column=col_idx, value=h_text)
        cell.font = FONT_HEADER
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER if col_idx in [3, 4] else ALIGN_LEFT
        cell.border = THIN_BORDER
    current_row += 1
    
    # Data Rows
    for row_data in table["columns"]:
        for col_idx, val in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = FONT_BODY
            cell.border = THIN_BORDER
            cell.alignment = ALIGN_CENTER if col_idx in [3, 4] else ALIGN_LEFT
        current_row += 1
        
    # Constraint row
    if table.get("constraint"):
        cell_c = ws.cell(row=current_row, column=1, value=table["constraint"])
        cell_c.font = FONT_CONSTRAINT
        current_row += 1
        
    current_row += 1  # Empty gap row between tables

# Adjust column widths
ws.column_dimensions["A"].width = 28
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 28
ws.column_dimensions["D"].width = 30

output_file = "ThietKe_CSDL_Nhom7_PhongHop_SinhNhat.xlsx"
wb.save(output_file)
print(f"Created {output_file} successfully!")
