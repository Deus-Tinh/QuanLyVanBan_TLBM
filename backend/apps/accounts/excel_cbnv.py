"""
Hỗ trợ xuất tệp mẫu và nhập dữ liệu Cán bộ Nhân viên (CBNV) bằng Excel (.xlsx).
Tác giả: Lê Duy Khánh (leduykhanhdevs).
"""
import re
import secrets
from io import BytesIO
from datetime import datetime, date
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.db import transaction
from django.contrib.auth import get_user_model

from organizations.models import DonVi
from .models import Employee, Position, vietnam_phone_regex


def generate_cbnv_template() -> BytesIO:
    """
    Sinh tệp Excel mẫu để nhập danh sách CBNV.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Danh Sach CBNV"

    # Header styling
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    hint_font = Font(name="Arial", size=9, italic=True, color="475569")
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    headers = [
        "Họ và tên (*)",
        "Ngày sinh (*)",
        "Email (*)",
        "Số điện thoại",
        "Tên Đơn vị / Phòng ban (*)",
        "Tên Chức vụ (*)",
        "Chức danh tự do",
        "Trạng thái (Đang làm việc / Nghỉ việc)",
        "Ghi chú",
    ]

    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        cell.border = thin_border

    # Sample rows
    sample_rows = [
        [
            "Nguyễn Văn An",
            "1990-05-15",
            "an.nguyen@baolong.vn",
            "0912345678",
            "Bảo Long Sài Gòn",
            "Trưởng phòng",
            "Trưởng phòng Kinh doanh",
            "Đang làm việc",
            "Nhân sự tuyển dụng đợt 1",
        ],
        [
            "Trần Thị Bích",
            "1994-11-20",
            "bich.tran@baolong.vn",
            "0987654321",
            "Ban Pháp chế",
            "Chuyên viên",
            "Chuyên viên Pháp lý",
            "Đang làm việc",
            "Nhân sự chính thức",
        ],
    ]

    # Hints row
    hints = [
        "(Ví dụ: Nguyễn Văn A)",
        "(Định dạng YYYY-MM-DD)",
        "(Ví dụ: an.nguyen@baolong.vn)",
        "(Ví dụ: 0912345678)",
        "(Tên đơn vị đã có trong hệ thống)",
        "(Tên chức vụ)",
        "(Chức danh nội bộ)",
        "(Đang làm việc / Nghỉ việc)",
        "(Ghi chú tự do)",
    ]
    ws.append(hints)
    for col_idx in range(1, len(hints) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = hint_font
        cell.alignment = align_center
        cell.border = thin_border

    for row_data in sample_rows:
        ws.append(row_data)

    # Set column widths
    col_widths = {
        "A": 26,  # Họ và tên
        "B": 16,  # Ngày sinh
        "C": 28,  # Email
        "D": 18,  # SĐT
        "E": 30,  # Đơn vị
        "F": 22,  # Chức vụ
        "G": 24,  # Chức danh
        "H": 25,  # Trạng thái
        "I": 30,  # Ghi chú
    }
    for col_letter, w in col_widths.items():
        ws.column_dimensions[col_letter].width = w

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _parse_date(val) -> date | None:
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    if not val:
        return None
    val_str = str(val).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(val_str, fmt).date()
        except ValueError:
            continue
    return None


def import_cbnv_from_excel(file_obj) -> dict:
    """
    Đọc và kiểm tra tính hợp lệ từng dòng của tệp Excel nhập danh sách CBNV.
    Trả về dict gồm:
      - success: bool
      - message: str
      - created_count: int
      - errors: list[dict(row=int, message=str)]
    """
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        ws = wb.active
    except Exception as e:
        return {
            "success": False,
            "message": f"Không thể đọc tệp Excel. Lỗi: {str(e)}",
            "errors": [{"row": 0, "message": "Tệp tải lên không phải định dạng Excel hợp lệ (.xlsx)."}],
        }

    rows = list(ws.iter_rows(values_only=True))
    if len(rows) <= 1:
        return {
            "success": False,
            "message": "Tệp Excel không có dòng dữ liệu nào ngoài dòng tiêu đề.",
            "errors": [{"row": 0, "message": "Tệp Excel trống."}],
        }

    # Tự động phát hiện có cột STT ở đầu hay không
    first_row = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
    offset = 1 if (len(first_row) > 0 and ("stt" in first_row[0] or "số tt" in first_row[0] or "số thứ tự" in first_row[0])) else 0

    errors = []
    valid_records = []
    seen_emails_in_file = set()

    for idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue  # Bỏ qua dòng trống hoàn toàn

        # Bỏ qua dòng hướng dẫn / gợi ý nếu ô bắt đầu bằng dấu ngoặc đơn '('
        first_cell = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
        second_cell = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if first_cell.startswith("(") or second_cell.startswith("("):
            continue

        def _get_val(col_idx):
            target = col_idx + offset
            if target < len(row) and row[target] is not None:
                return row[target]
            return None

        ho_ten = str(_get_val(0)).strip() if _get_val(0) is not None else ""
        raw_birthday = _get_val(1)
        raw_email = str(_get_val(2)).strip().lower() if _get_val(2) is not None else ""
        sdt = str(_get_val(3)).strip() if _get_val(3) is not None else ""
        don_vi_name = str(_get_val(4)).strip() if _get_val(4) is not None else ""
        chuc_vu_name = str(_get_val(5)).strip() if _get_val(5) is not None else ""
        chuc_danh = str(_get_val(6)).strip() if _get_val(6) is not None else ""
        raw_trang_thai = str(_get_val(7)).strip() if _get_val(7) is not None else "Đang làm việc"
        ghi_chu = str(_get_val(8)).strip() if _get_val(8) is not None else ""

        # 1. Validate Họ và tên
        if not ho_ten:
            errors.append({"row": idx, "message": "Họ và tên không được để trống."})

        # 2. Validate Ngày sinh
        birthday = _parse_date(raw_birthday)
        if not birthday:
            errors.append({
                "row": idx,
                "message": f"Ngày sinh '{raw_birthday}' không đúng định dạng (hỗ trợ YYYY-MM-DD hoặc DD/MM/YYYY).",
            })

        # 3. Validate Email
        if not raw_email:
            errors.append({"row": idx, "message": "Email không được để trống."})
        elif not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", raw_email):
            errors.append({"row": idx, "message": f"Email '{raw_email}' không đúng định dạng."})
        elif raw_email in seen_emails_in_file:
            errors.append({"row": idx, "message": f"Email '{raw_email}' bị trùng lặp trong chính tệp Excel tải lên."})
        elif get_user_model().objects.filter(email__iexact=raw_email).exists():
            errors.append({"row": idx, "message": f"Email '{raw_email}' đã tồn tại trên hệ thống."})
        else:
            seen_emails_in_file.add(raw_email)

        # 4. Validate Số điện thoại
        if sdt:
            sdt_clean = sdt.replace(" ", "").replace(".", "")
            if not re.match(r"^(0|\+84)[1-9][0-9]{7,9}$", sdt_clean):
                errors.append({"row": idx, "message": f"Số điện thoại '{sdt}' không hợp lệ theo chuẩn Việt Nam."})
        else:
            sdt_clean = ""

        # 5. Validate Đơn vị / Phòng ban
        don_vi_instance = None
        if not don_vi_name:
            errors.append({"row": idx, "message": "Tên Đơn vị / Phòng ban không được để trống."})
        else:
            don_vi_instance = DonVi.objects.filter(ten__iexact=don_vi_name).first()
            if not don_vi_instance:
                errors.append({"row": idx, "message": f"Không tìm thấy Đơn vị có tên '{don_vi_name}' trong hệ thống."})

        # 6. Validate Chức vụ
        chuc_vu_instance = None
        if not chuc_vu_name:
            errors.append({"row": idx, "message": "Tên Chức vụ không được để trống."})
        else:
            chuc_vu_instance = Position.objects.filter(name__iexact=chuc_vu_name).first()
            if not chuc_vu_instance:
                # Nếu chưa có, tạo tự động chức vụ mới
                next_order = Position.get_next_display_order()
                chuc_vu_instance = Position.objects.create(name=chuc_vu_name, display_order=next_order)

        # 7. Validate Trạng thái
        trang_thai = Employee.Status.LAM_VIEC
        if "nghỉ" in raw_trang_thai.lower() or "nghi" in raw_trang_thai.lower():
            trang_thai = Employee.Status.NGHI_VIEC

        valid_records.append({
            "ho_ten": ho_ten,
            "ngay_sinh": birthday,
            "email": raw_email,
            "sdt": sdt_clean,
            "don_vi": don_vi_instance,
            "chuc_vu": chuc_vu_instance,
            "chuc_danh": chuc_danh,
            "trang_thai": trang_thai,
            "ghi_chu": ghi_chu,
        })

    if errors:
        return {
            "success": False,
            "errors": errors,
            "message": f"Phát hiện {len(errors)} lỗi trong tệp Excel. Vui lòng kiểm tra và sửa lại trước khi nhập.",
        }

    # Nếu toàn bộ dữ liệu hợp lệ, tiến hành lưu nguyên tử (atomic) vào CSDL
    created_count = 0
    UserModel = get_user_model()
    with transaction.atomic():
        for item in valid_records:
            email = item["email"]
            local_part = email.split("@")[0]
            candidate = local_part
            idx = 1
            while UserModel.objects.filter(username__iexact=candidate).exists():
                candidate = f"{local_part}{idx}"
                idx += 1

            UserModel.objects.create_user(
                username=candidate,
                password=f"ITC@{secrets.token_hex(4)}",
                email=email,
                ho_ten=item["ho_ten"],
                ngay_sinh=item["ngay_sinh"],
                don_vi=item["don_vi"],
                chuc_vu=item["chuc_vu"],
                chuc_danh=item["chuc_danh"],
                sdt=item["sdt"],
                trang_thai=item["trang_thai"],
                is_active=(item["trang_thai"] != Employee.Status.NGHI_VIEC),
                ghi_chu=item["ghi_chu"],
            )
            created_count += 1

    return {
        "success": True,
        "message": f"Nhập thành công {created_count} cán bộ nhân viên.",
        "created_count": created_count,
        "errors": [],
    }
