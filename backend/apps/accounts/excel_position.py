"""
Tệp mẫu và import danh mục Chức vụ bằng Excel (.xlsx) — Issue #20.

Import kiểm tra từng dòng và trả lỗi cụ thể theo dòng (row = số dòng trong file Excel).
Dòng hợp lệ được lưu; dòng lỗi bị bỏ qua và báo lại, không dừng cả file.
"""
from io import BytesIO

import openpyxl
from django.db import IntegrityError, transaction
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import MAX_DISPLAY_ORDER, Position

HEADERS = ["STT ưu tiên", "Tên chức vụ (*)", "Hiển thị công khai (Có/Không)", "Ghi chú"]
MAX_ROWS = 1000
MAX_FILE_SIZE = 2 * 1024 * 1024
TRUE_VALUES = {"co", "có", "true", "1", "yes", "y", "x", "cong khai", "công khai", "public"}
FALSE_VALUES = {"khong", "không", "false", "0", "no", "n", "noi bo", "nội bộ", "internal"}

ERR_INVALID_FILE = "positions.import.errors.invalidFile"
ERR_INVALID_TEMPLATE = "positions.import.errors.invalidTemplate"
ERR_EMPTY = "positions.import.errors.emptyFile"
ERR_TOO_MANY = "positions.import.errors.tooManyRows"
ERR_TOO_LARGE = "positions.import.errors.fileTooLarge"


class ImportFileError(Exception):
    """Lỗi cả file (không đọc được, sai mẫu, rỗng...): API trả 400 kèm key i18n."""

    def __init__(self, key: str):
        super().__init__(key)
        self.key = key


def _normalize_header(value) -> str:
    text = str(value or "").replace("(*)", "")
    return " ".join(text.split()).casefold()


def generate_position_template() -> BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Chuc vu"

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    side = Side(style="thin", color="CBD5E1")
    border = Border(left=side, right=side, top=side, bottom=side)

    ws.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font, cell.alignment, cell.border = header_fill, header_font, center, border
    for letter, width in {"A": 16, "B": 36, "C": 30, "D": 40}.items():
        ws.column_dimensions[letter].width = width

    guide = wb.create_sheet("Huong dan")
    guide_rows = [
        ["Hướng dẫn nhập danh mục Chức vụ"],
        ["Nhập dữ liệu ở sheet 'Chuc vu' từ dòng 2, không đổi tên các cột dòng 1."],
        ["STT ưu tiên: số nguyên dương, không trùng. Để trống thì hệ thống tự lấy số lớn nhất + 1."],
        ["Tên chức vụ (*): bắt buộc, không trùng tên đã có (không phân biệt hoa/thường)."],
        ["Hiển thị công khai: Có hoặc Không. Để trống là Có."],
        [f"Tối đa {MAX_ROWS} dòng dữ liệu, dung lượng tối đa {MAX_FILE_SIZE // (1024 * 1024)} MB."],
        ["Dòng lỗi sẽ được báo rõ số dòng và lý do; các dòng hợp lệ vẫn được lưu."],
    ]
    for row in guide_rows:
        guide.append(row)
    guide["A1"].font = Font(name="Arial", size=12, bold=True)
    guide.column_dimensions["A"].width = 100

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _parse_stt(raw):
    """Trả về (giá trị | None, mã lỗi | None). None kèm không lỗi nghĩa là để hệ thống tự cấp."""
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return None, None
    if isinstance(raw, bool):
        return None, "positions.import.errors.invalidStt"
    if isinstance(raw, float):
        if raw != int(raw):
            return None, "positions.import.errors.invalidStt"
        raw = int(raw)
    if isinstance(raw, str):
        try:
            raw = int(raw.strip())
        except ValueError:
            return None, "positions.import.errors.invalidStt"
    if not isinstance(raw, int) or raw < 1:
        return None, "positions.import.errors.invalidStt"
    if raw > MAX_DISPLAY_ORDER:
        return None, "positions.messages.sttTooLarge"
    return raw, None


def _parse_public(raw):
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        return True, None
    if isinstance(raw, bool):
        return raw, None
    text = str(raw).strip().casefold()
    if text in TRUE_VALUES:
        return True, None
    if text in FALSE_VALUES:
        return False, None
    return True, "positions.import.errors.invalidPublic"


def import_positions_from_excel(file_obj) -> dict:
    """
    Đọc file .xlsx và lưu các dòng hợp lệ.
    Trả về {"total", "imported", "failed", "errors": [{"row", "field", "code", "value"}]}.
    """
    if getattr(file_obj, "size", 0) > MAX_FILE_SIZE:
        raise ImportFileError(ERR_TOO_LARGE)

    try:
        workbook = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        sheet = workbook.worksheets[0]
        rows = list(sheet.iter_rows(values_only=True))
    except Exception as exc:  # openpyxl ném nhiều loại lỗi khác nhau với file hỏng/sai định dạng
        raise ImportFileError(ERR_INVALID_FILE) from exc

    if not rows:
        raise ImportFileError(ERR_EMPTY)

    header = [_normalize_header(c) for c in rows[0][: len(HEADERS)]]
    if header != [_normalize_header(h) for h in HEADERS]:
        raise ImportFileError(ERR_INVALID_TEMPLATE)

    data_rows = [
        (index, tuple(row[: len(HEADERS)]) + (None,) * (len(HEADERS) - len(row)))
        for index, row in enumerate(rows[1:], start=2)
        if any(cell is not None and str(cell).strip() != "" for cell in row)
    ]
    if not data_rows:
        raise ImportFileError(ERR_EMPTY)
    if len(data_rows) > MAX_ROWS:
        raise ImportFileError(ERR_TOO_MANY)

    errors = []
    candidates = []  # (row, stt | None, name, is_public, note)
    seen_names = {}
    seen_stt = {}
    existing_names = {n.casefold() for n in Position.objects.values_list("name", flat=True)}
    existing_stt = set(Position.objects.values_list("display_order", flat=True))

    def add_error(row, field, code, value=None):
        errors.append({"row": row, "field": field, "code": code, "value": "" if value is None else str(value)})

    for row_number, (stt_raw, name_raw, public_raw, note_raw) in data_rows:
        row_errors = 0

        name = " ".join(str(name_raw or "").split())
        if not name:
            add_error(row_number, "ten_chuc_vu", "positions.messages.nameRequired")
            row_errors += 1
        elif len(name) > 255:
            add_error(row_number, "ten_chuc_vu", "positions.messages.nameTooLong", name[:40])
            row_errors += 1
        elif name.casefold() in existing_names:
            add_error(row_number, "ten_chuc_vu", "positions.messages.nameExists", name)
            row_errors += 1
        elif name.casefold() in seen_names:
            add_error(row_number, "ten_chuc_vu", "positions.import.errors.duplicateNameInFile", name)
            row_errors += 1

        stt, stt_error = _parse_stt(stt_raw)
        if stt_error:
            add_error(row_number, "stt_uu_tien", stt_error, stt_raw)
            row_errors += 1
        elif stt is not None and stt in existing_stt:
            add_error(row_number, "stt_uu_tien", "positions.messages.sttExists", stt)
            row_errors += 1
        elif stt is not None and stt in seen_stt:
            add_error(row_number, "stt_uu_tien", "positions.import.errors.duplicateSttInFile", stt)
            row_errors += 1

        is_public, public_error = _parse_public(public_raw)
        if public_error:
            add_error(row_number, "is_public", public_error, public_raw)
            row_errors += 1

        if row_errors:
            continue

        seen_names[name.casefold()] = row_number
        if stt is not None:
            seen_stt[stt] = row_number
        candidates.append((row_number, stt, name, is_public, str(note_raw or "").strip()))

    # Dòng không có STT lấy số lớn nhất + 1, tránh cả các STT khai báo tường minh trong file.
    reserved = existing_stt | set(seen_stt)
    next_stt = max(reserved, default=0) + 1
    if next_stt > MAX_DISPLAY_ORDER:
        # STT lớn nhất đã chạm trần cột INTEGER: bắt đầu tìm số trống từ 1 (vòng while bên dưới bỏ qua số đã dùng).
        next_stt = 1
    imported = 0
    for row_number, stt, name, is_public, note in candidates:
        if stt is None:
            while next_stt in reserved:
                next_stt += 1
            stt = next_stt
            reserved.add(stt)
        try:
            with transaction.atomic():
                Position.objects.create(display_order=stt, name=name, is_public=is_public, note=note)
            imported += 1
        except IntegrityError:
            add_error(row_number, "stt_uu_tien", "positions.messages.sttExists", stt)

    errors.sort(key=lambda e: e["row"])
    return {
        "total": len(data_rows),
        "imported": imported,
        "failed": len({e["row"] for e in errors}),
        "errors": errors,
    }
