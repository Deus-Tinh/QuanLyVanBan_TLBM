from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError


REQUIRED_COLUMNS = {
    "CBNV": {"Tên đăng nhập", "Họ và tên", "Ngày sinh", "Khối", "Đơn vị", "Chức vụ"},
    "CBNV_BoSung": {
        "Email (khóa nối với CBNV_Import)",
        "Tên đăng nhập",
        "Nhóm quyền",
    },
    "PhongHop": {"Tên phòng", "Vị trí", "Sức chứa"},
    "DatPhongHop": {"Phòng", "Người đặt", "Ngày (mẫu)", "Bắt đầu", "Kết thúc"},
    "LoiChucSinhNhat": {"Người nhận", "Người gửi", "Nội dung"},
}


class Command(BaseCommand):
    help = "Kiểm tra workbook dữ liệu mẫu"

    def add_arguments(self, parser):
        parser.add_argument("--file", required=True)

    def handle(self, *args, **options):
        file_path = Path(options["file"]).expanduser()

        if not file_path.is_file():
            raise CommandError(f"Không tìm thấy file: {file_path}")

        try:
            workbook = openpyxl.load_workbook(
                file_path,
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise CommandError(f"Không đọc được file Excel: {exc}") from exc

        try:
            missing_sheets = set(REQUIRED_COLUMNS) - set(workbook.sheetnames)
            if missing_sheets:
                raise CommandError(
                    f"Thiếu sheet: {', '.join(sorted(missing_sheets))}"
                )

            for sheet_name, required_columns in REQUIRED_COLUMNS.items():
                rows = workbook[sheet_name].iter_rows(values_only=True)
                header_row = next(rows, None)

                if header_row is None:
                    raise CommandError(f"Sheet {sheet_name} đang trống.")

                headers = {
                    str(value).strip()
                    for value in header_row
                    if value is not None
                }
                missing_columns = required_columns - headers

                if missing_columns:
                    raise CommandError(
                        f"Sheet {sheet_name} thiếu cột: "
                        f"{', '.join(sorted(missing_columns))}"
                    )

                count = sum(
                    1 for row in rows
                    if any(value is not None for value in row)
                )
                self.stdout.write(f"{sheet_name}: {count} dòng")

        finally:
            workbook.close()

        self.stdout.write("Kiểm tra xong; chưa ghi dữ liệu vào DB.")
