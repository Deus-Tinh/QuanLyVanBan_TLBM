import io
import openpyxl
from django.core.exceptions import ValidationError
from django.test import TestCase
from rest_framework.test import APIClient
from rest_framework import status
from django.contrib.auth.models import Permission, User

from organizations.models import KhoiKhuVuc, DonVi


class OrganizationModelAndAPITest(TestCase):
    """
    Test suite kiểm tra phân hệ Organizations: Model, Validation Keys, và API Endpoints.
    """

    def setUp(self):
        self.user = User.objects.create_user(
            username="testuser",
            password="testpassword123",
            email="test@itc.edu.vn",
        )
        self.export_permission = Permission.objects.get(
            content_type__app_label="organizations",
            codename="export_donvi",
        )
        self.user.user_permissions.add(self.export_permission)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)

        self.region = KhoiKhuVuc.objects.create(ten="Khối Cơ quan")
        self.unit = DonVi.objects.create(
            khoi=self.region,
            stt_uu_tien=1,
            ten="Phòng Kế toán",
            dien_thoai="02838123456",
            email="ketoan@itc.edu.vn",
        )

    def test_khoi_khuvuc_creation(self):
        self.assertEqual(str(self.region), "Khối Cơ quan")
        self.assertEqual(self.region.don_vi_list.count(), 1)

    def test_donvi_creation(self):
        self.assertEqual(self.unit.stt_uu_tien, 1)
        self.assertTrue(self.unit.is_public)

    def test_auto_priority_assignment(self):
        data = {
            "khoi": self.region.id,
            "ten": "Phòng Tổ chức Cán bộ",
            "dien_thoai": "0912345678",
            "email": "tccb@itc.edu.vn",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["stt_uu_tien"], 2)

    def test_validation_key_on_blank_name(self):
        data = {
            "khoi": self.region.id,
            "ten": "   ",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.nameRequired", str(response.data))

    def test_model_save_rejects_invalid_values(self):
        """Kiểm tra model tự chặn dữ liệu sai khi lưu trực tiếp, không qua API."""
        with self.assertRaises(ValidationError) as region_context:
            KhoiKhuVuc(ten="   ").save()
        self.assertIn("organizations.messages.nameRequired", str(region_context.exception))

        with self.assertRaises(ValidationError) as unit_context:
            DonVi(
                khoi=self.region,
                stt_uu_tien=2,
                ten="Đơn vị không hợp lệ",
                dien_thoai="12345678",
            ).save()
        self.assertIn("organizations.messages.phoneInvalid", str(unit_context.exception))

    def test_prevent_delete_region_with_units(self):
        response = self.client.delete(f"/api/organizations/regions/{self.region.id}/")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertEqual(response.data.get("key"), "organizations.messages.regionHasUnits")

    def test_export_excel_endpoint_and_content(self):
        """Kiểm tra tệp Excel được sinh ra với tiêu đề song ngữ và nội dung chính xác."""
        # 1. Kiểm tra xuất tiếng Việt
        res_vi = self.client.get("/api/organizations/units/export/?lang=vi")
        self.assertEqual(res_vi.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res_vi["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        wb_vi = openpyxl.load_workbook(io.BytesIO(res_vi.content))
        ws_vi = wb_vi.active
        self.assertEqual(ws_vi.title, "Danh Sách Đơn Vị")
        headers_vi = [cell.value for cell in ws_vi[1]]
        self.assertIn("STT Ưu Tiên", headers_vi)
        self.assertIn("Tên Đơn Vị / Phòng Ban", headers_vi)
        # Kiểm tra nội dung hàng 2
        row2_vi = [cell.value for cell in ws_vi[2]]
        self.assertEqual(row2_vi[1], "Phòng Kế toán")

        # 2. Kiểm tra xuất tiếng Anh
        res_en = self.client.get("/api/organizations/units/export/?lang=en")
        self.assertEqual(res_en.status_code, status.HTTP_200_OK)
        wb_en = openpyxl.load_workbook(io.BytesIO(res_en.content))
        ws_en = wb_en.active
        self.assertEqual(ws_en.title, "Departments Directory")
        headers_en = [cell.value for cell in ws_en[1]]
        self.assertIn("Priority Index", headers_en)
        self.assertIn("Unit / Department Name", headers_en)

    def test_auto_priority_on_blank_string(self):
        """Kiểm tra khi frontend gửi chuỗi rỗng stt_uu_tien="", hệ thống vẫn tự động gán max + 1."""
        data = {
            "khoi": self.region.id,
            "stt_uu_tien": "",
            "ten": "Đơn Vị STT Chuỗi Rỗng",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["stt_uu_tien"], 2)

    def test_export_requires_permission(self):
        """Kiểm tra người dùng đã đăng nhập nhưng thiếu quyền export bị từ chối."""
        self.user.user_permissions.remove(self.export_permission)

        response = self.client.get("/api/organizations/units/export/")

        self.assertEqual(response.status_code, status.HTTP_403_FORBIDDEN)
        self.assertEqual(response.data["detail"], "organizations.messages.exportPermissionDenied")

    def test_duplicate_priority_rejection(self):
        """Kiểm tra không cho phép trùng STT ưu tiên đã tồn tại."""
        data = {
            "khoi": self.region.id,
            "stt_uu_tien": 1,
            "ten": "Đơn vị Trùng STT",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.priorityExists", str(response.data))

    def test_invalid_phone_number_rejection(self):
        """Kiểm tra regex số điện thoại Việt Nam hợp lệ."""
        data = {
            "khoi": self.region.id,
            "ten": "Đơn vị SĐT Lỗi",
            "dien_thoai": "12345678",  # Không bắt đầu bằng 0
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.phoneInvalid", str(response.data))

    def test_invalid_email_format(self):
        """Kiểm tra định dạng email không hợp lệ."""
        data = {
            "khoi": self.region.id,
            "ten": "Đơn vị Email Lỗi",
            "email": "not-a-valid-email",
        }
        response = self.client.post("/api/organizations/units/", data, format="json")
        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("organizations.messages.emailInvalid", str(response.data))

    def test_region_pagination_and_all_param(self):
        """Kiểm tra phân trang mặc định nhiều trang và tham số ?all=true cho Khối/Khu vực."""
        # Tạo thêm 12 Khối để có tổng cộng 13 Khối (vượt page_size = 10)
        KhoiKhuVuc.objects.bulk_create([
            KhoiKhuVuc(ten=f"Khối Thử Nghiệm Số {i:02d}") for i in range(2, 14)
        ])

        # 1. Gọi trang 1 -> phải có 10 bản ghi và có liên kết trang tiếp theo (next)
        res_page1 = self.client.get("/api/organizations/regions/?page=1")
        self.assertEqual(res_page1.status_code, status.HTTP_200_OK)
        self.assertEqual(res_page1.data["count"], 13)
        self.assertEqual(len(res_page1.data["results"]), 10)
        self.assertIsNotNone(res_page1.data["next"])

        # 2. Gọi trang 2 -> phải có 3 bản ghi còn lại và có liên kết trang trước (previous)
        res_page2 = self.client.get("/api/organizations/regions/?page=2")
        self.assertEqual(res_page2.status_code, status.HTTP_200_OK)
        self.assertEqual(len(res_page2.data["results"]), 3)
        self.assertIsNotNone(res_page2.data["previous"])

        # 3. Gọi với ?all=true -> phải trả về toàn bộ 13 bản ghi dạng mảng không phân trang
        res_all = self.client.get("/api/organizations/regions/?all=true")
        self.assertEqual(res_all.status_code, status.HTTP_200_OK)
        self.assertIsInstance(res_all.data, list)
        self.assertEqual(len(res_all.data), 13)

    def test_unit_search_by_address(self):
        """Kiểm tra tìm kiếm đơn vị theo địa chỉ."""
        self.unit.dia_chi = "254 Nguyễn Văn Linh, Đà Nẵng"
        self.unit.save()

        response = self.client.get("/api/organizations/units/?search=Nguyễn+Văn+Linh")
        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["count"], 1)

        response_not_found = self.client.get("/api/organizations/units/?search=Hà+Nội")
        self.assertEqual(response_not_found.status_code, status.HTTP_200_OK)
        self.assertEqual(response_not_found.data["count"], 0)

    def test_export_search_by_address(self):
        """Kiểm tra xuất file Excel lọc đúng địa chỉ và loại trừ các bản ghi không khớp."""
        # Đơn vị 1 khớp từ khóa Đà Nẵng
        self.unit.dia_chi = "254 Nguyễn Văn Linh, Đà Nẵng"
        self.unit.save()

        # Tạo Đơn vị 2 có địa chỉ ở Hà Nội (không khớp từ khóa tìm kiếm)
        DonVi.objects.create(
            khoi=self.region,
            stt_uu_tien=99,
            ten="Văn Phòng Đại Diện",
            dia_chi="123 Phố Cầu Giấy, Hà Nội",
        )

        # 1. Xuất file với từ khóa 'Nguyễn Văn Linh'
        response = self.client.get("/api/organizations/units/export/?search=Nguyễn+Văn+Linh")
        self.assertEqual(response.status_code, status.HTTP_200_OK)

        # Mở file Excel từ luồng byte và đọc các dòng
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))

        # Phải có đúng 2 dòng: dòng 1 là tiêu đề cột, dòng 2 là đơn vị khớp địa chỉ
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][1], "Phòng Kế toán")
        self.assertIn("Nguyễn Văn Linh", rows[1][5])

        # 2. Xuất file với từ khóa không tồn tại -> chỉ có đúng 1 dòng tiêu đề cột
        res_empty = self.client.get("/api/organizations/units/export/?search=Địa+Chỉ+Không+Tồn+Tại")
        wb_empty = openpyxl.load_workbook(io.BytesIO(res_empty.content))
        ws_empty = wb_empty.active
        rows_empty = list(ws_empty.iter_rows(values_only=True))
        self.assertEqual(len(rows_empty), 1)

    def test_unit_update_and_delete(self):
        """Kiểm tra cập nhật và xóa đơn vị thành công."""
        # Cập nhật PUT
        update_data = {
            "khoi": self.region.id,
            "stt_uu_tien": 1,
            "ten": "Phòng Kế toán Đổi Tên",
            "dien_thoai": "02838999999",
            "email": "ketoan_new@itc.edu.vn",
            "is_public": False,
        }
        res_put = self.client.put(f"/api/organizations/units/{self.unit.id}/", update_data, format="json")
        self.assertEqual(res_put.status_code, status.HTTP_200_OK)
        self.assertEqual(res_put.data["ten"], "Phòng Kế toán Đổi Tên")
        self.assertFalse(res_put.data["is_public"])

        # Xóa DELETE
        res_delete = self.client.delete(f"/api/organizations/units/{self.unit.id}/")
        self.assertEqual(res_delete.status_code, status.HTTP_204_NO_CONTENT)
        self.assertFalse(DonVi.objects.filter(pk=self.unit.id).exists())

    def test_unauthenticated_access_denied(self):
        """Kiểm tra bắt buộc xác thực đăng nhập (401 Unauthorized)."""
        unauthenticated_client = APIClient()
        response = unauthenticated_client.get("/api/organizations/units/")
        self.assertEqual(response.status_code, status.HTTP_401_UNAUTHORIZED)
