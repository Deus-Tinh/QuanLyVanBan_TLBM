"""Tạo bộ tài liệu mẫu phong phú ở nhiều thư mục khác nhau để kiểm thử Trợ lý AI TLBM

Hỗ trợ cả định dạng DOCX và TXT với nội dung quy chuẩn văn bản hành chính doanh nghiệp.
"""

import json
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path

# Đảm bảo in UTF-8 không bị lỗi trên Windows console (cp1258)
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

ROOT_DIR = Path(__file__).resolve().parent.parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Thư mục đích chứa các thư mục con phân loại tài liệu kiểm thử
TEST_DOCS_ROOT = ROOT_DIR / "tai_lieu_test_ai"


DOCUMENTS_DATA = [
    {
        "folder": "01_Bao_Mat_Va_Luu_Tru",
        "category": "Bảo mật & Lưu trữ",
        "filename_base": "Quy_Che_Bao_Mat_Thong_Tin_TLBM",
        "title": "QUY CHẾ BẢO MẬT THÔNG TIN VÀ QUẢN LÝ TÀI LIỆU BẢO MẬT NỘI BỘ",
        "so_hieu": "18/2026/QC-TLBM",
        "sections": [
            {
                "heading": "Chương I: QUY ĐỊNH CHUNG VÀ CÁC CẤP ĐỘ BẢO MẬT",
                "content": (
                    "Điều 1. Phạm vi điều chỉnh và đối tượng áp dụng\n"
                    "Quy chế này quy định việc phân loại, xác định độ mật, lưu trữ, tiếp nhận, chuyển giao, "
                    "và hủy bỏ các tài liệu mang tính chất bí mật kinh doanh, bí mật kỹ thuật và công văn nội bộ "
                    "trong toàn bộ tổ chức TLBM. Toàn thể cán bộ, nhân viên, thực tập sinh và cộng tác viên "
                    "đều có nghĩa vụ tuân thủ nghiêm ngặt quy chế này.\n\n"
                    "Điều 2. Phân loại 3 cấp độ bảo mật tài liệu\n"
                    "1. Cấp độ 1 - TUYỆT MẬT (Top Secret): Gồm mã khóa mật mã hệ thống, cấu hình chứng thư số gốc (Root CA), "
                    "kế hoạch sáp nhập tài chính và chiến lược kinh doanh cấp hội đồng quản trị. "
                    "Chỉ có Chủ tịch HĐQT, Tổng Giám Đốc và Quản trị trưởng hệ thống mới có quyền truy cập.\n"
                    "2. Cấp độ 2 - TỐI MẬT (Secret): Gồm hồ sơ dự thầu đang triển khai, báo cáo kiểm toán tài chính trước phát hành, "
                    "và mã nguồn các module nghiệp vụ lõi (Signing Engine, Security Barrier Views).\n"
                    "3. Cấp độ 3 - MẬT (Confidential): Gồm danh sách lương và nhân sự chi tiết, biên bản họp giao ban phòng ban, "
                    "và hợp đồng kinh tế đang trong giai đoạn thương thảo."
                ),
            },
            {
                "heading": "Chương II: THỜI HẠN BẢO QUẢN VÀ CHẾ TÀI XỬ PHẠT",
                "content": (
                    "Điều 3. Thời hạn bảo quản tài liệu mật\n"
                    "- Tài liệu Tuyệt mật: Lưu trữ vĩnh viễn trong kho lưu trữ dữ liệu mã hóa chuẩn AES-256 có phân tán khóa.\n"
                    "- Tài liệu Tối mật: Thời hạn lưu trữ tối thiểu là 20 năm tính từ ngày văn bản có hiệu lực.\n"
                    "- Tài liệu Mật: Thời hạn lưu trữ tối thiểu là 10 năm tính từ ngày văn bản có hiệu lực.\n\n"
                    "Điều 4. Chế tài xử phạt hành vi vi phạm bảo mật\n"
                    "Cán bộ, nhân viên có hành vi sao chép trái phép, chụp ảnh màn hình văn bản mật, "
                    "hoặc để lộ lọt thông tin bảo mật ra bên ngoài sẽ bị xử lý kỷ luật theo các mức độ:\n"
                    "- Khiển trách và đình chỉ công tác 15 ngày đối với vi phạm lần đầu không cố ý.\n"
                    "- Phạt tiền từ 20.000.000 VNĐ đến 50.000.000 VNĐ để bồi thường thiệt hại.\n"
                    "- Chấm dứt hợp đồng lao động ngay lập tức và chuyển hồ sơ sang cơ quan công an xử lý hình sự nếu gây hậu quả nghiêm trọng."
                ),
            },
        ],
    },
    {
        "folder": "02_Quy_Trinh_Trinh_Ky_Dien_Tu",
        "category": "Quy trình Trình ký",
        "filename_base": "Quy_Trinh_Ky_So_Va_Trinh_Ky_Van_Ban",
        "title": "QUY TRÌNH TRÌNH KÝ ĐIỆN TỬ VÀ KÝ SỐ VĂN BẢN TRÊN HỆ THỐNG TLBM",
        "so_hieu": "35/2026/QĐ-TLBM",
        "sections": [
            {
                "heading": "Chương I: LUỒNG DUYỆT VĂN BẢN ĐIỆN TỬ 4 BƯỚC",
                "content": (
                    "Điều 1. Các bước thực hiện trình ký điện tử\n"
                    "Quy trình xử lý hồ sơ văn bản điện tử bao gồm đúng 4 bước tiêu chuẩn:\n"
                    "• Bước 1 - Soạn thảo và đính kèm: Chuyên viên nghiệp vụ soạn thảo văn bản, đính kèm tệp văn bản chuẩn (.docx hoặc .pdf) "
                    "và tạo phiếu trình ký trên phần mềm TLBM, chọn danh sách người phối hợp và người duyệt.\n"
                    "• Bước 2 - Thẩm định và ký nháy: Trưởng phòng chuyên môn kiểm tra tính pháp lý, thể thức và nội dung. "
                    "Nếu hợp lệ, Trưởng phòng thực hiện ký nháy điện tử trong vòng 24 giờ làm việc kể từ lúc nhận phiếu.\n"
                    "• Bước 3 - Phê duyệt và ký số chính thức: Lãnh đạo đơn vị (Ban Giám Đốc) xem xét toàn bộ hồ sơ và thực hiện ký số "
                    "bằng USB Token hoặc Chữ ký số từ xa (Remote Signing). Hệ thống yêu cầu xác thực mã PIN bảo mật cá nhân.\n"
                    "• Bước 4 - Phát hành và lưu trữ: Văn thư cơ quan kiểm tra tính nguyên vẹn của chữ ký số, lấy số hiệu chính thức, "
                    "đóng dấu số điện tử của cơ quan và vào sổ công văn phát hành."
                ),
            },
            {
                "heading": "Chương II: THỜI HẠN XỬ LÝ VÀ TIÊU CHUẨN KỸ THUẬT KÝ SỐ",
                "content": (
                    "Điều 2. Thời hạn xử lý phiếu trình ký\n"
                    "- Văn bản thông thường: Toàn bộ quá trình từ lúc lập phiếu đến khi hoàn tất ký số không quá 48 giờ làm việc.\n"
                    "- Văn bản Khẩn: Thời hạn xử lý tối đa là 12 giờ làm việc.\n"
                    "- Văn bản Hỏa tốc / Thượng khẩn: Xử lý ngay lập tức, tối đa không quá 04 giờ làm việc.\n\n"
                    "Điều 3. Tiêu chuẩn kỹ thuật chữ ký số\n"
                    "Chữ ký số hợp lệ trên hệ thống TLBM phải đáp ứng chuẩn mật mã RSA 2048-bit hoặc ECDSA P-256, "
                    "hàm băm an toàn SHA-256, và có chứng thư số số hóa được cấp bởi các tổ chức chứng thực CA quốc gia hợp pháp."
                ),
            },
        ],
    },
    {
        "folder": "03_Che_Do_Va_Phuc_Loi_Nhan_Su",
        "category": "Nhân sự & Phúc lợi",
        "filename_base": "Chinh_Sach_Phuc_Loi_Va_Nghi_Phep_2026",
        "title": "QUY ĐỊNH VỀ CHẾ ĐỘ NGHỈ PHÉP, LÀM THÊM GIỜ VÀ PHÚC LỢI NHÂN SỰ NĂM 2026",
        "so_hieu": "52/2026/QĐ-NS",
        "sections": [
            {
                "heading": "Chương I: THỜI GIAN LÀM VIỆC VÀ NGHỈ PHÉP NĂM",
                "content": (
                    "Điều 1. Khung giờ làm việc tiêu chuẩn\n"
                    "- Thời gian làm việc: Từ Thứ Hai đến hết Thứ Sáu hàng tuần (Thứ Bảy và Chủ Nhật nghỉ).\n"
                    "- Buổi sáng: Từ 08:00 đến 12:00.\n"
                    "- Buổi chiều: Từ 13:30 đến 17:30.\n"
                    "- Thời gian nghỉ trưa: 1 giờ 30 phút (từ 12:00 đến 13:30).\n\n"
                    "Điều 2. Chế độ nghỉ phép năm (Annual Leave)\n"
                    "1. Nhân viên ký hợp đồng chính thức đủ 12 tháng được hưởng 14 ngày nghỉ phép năm có hưởng nguyên lương.\n"
                    "2. Thâm niên: Cứ làm việc đủ 03 năm liên tục tại cơ quan được cộng thêm 01 ngày phép năm vào quỹ phép.\n"
                    "3. Phép năm chưa sử dụng hết được phép chuyển tối đa 05 ngày sang quý I của năm kế tiếp."
                ),
            },
            {
                "heading": "Chương II: TIỀN LƯƠNG LÀM THÊM GIỜ (OT) VÀ PHỤ CẤP PHÚC LỢI",
                "content": (
                    "Điều 3. Chế độ tính lương làm thêm giờ (Overtime - OT)\n"
                    "Khi có lệnh huy động làm thêm giờ từ Trưởng bộ phận:\n"
                    "- OT vào ngày làm việc thường: Hưởng 150% tiền lương theo giờ.\n"
                    "- OT vào ngày nghỉ hàng tuần (Thứ Bảy, Chủ Nhật): Hưởng 200% tiền lương theo giờ.\n"
                    "- OT vào ngày lễ, Tết, ngày nghỉ có hưởng lương: Hưởng 300% tiền lương theo giờ (chưa kể tiền lương ngày lễ).\n\n"
                    "Điều 4. Các khoản phụ cấp và phúc lợi hàng tháng\n"
                    "- Phụ cấp ăn trưa: 750.000 VNĐ / nhân sự / tháng (chi trả vào kỳ lương đầu tháng).\n"
                    "- Phụ cấp xăng xe và điện thoại: 500.000 VNĐ / tháng đối với chuyên viên; 1.200.000 VNĐ / tháng đối với Trưởng phòng và Lãnh đạo.\n"
                    "- Khám sức khỏe định kỳ: Tổ chức vào tháng 06 hàng năm tại bệnh viện quốc tế với gói tiêu chuẩn 4.500.000 VNĐ / người."
                ),
            },
        ],
    },
    {
        "folder": "04_Huong_Dan_Ky_Thuat_CNTT",
        "category": "Kỹ thuật CNTT",
        "filename_base": "Huong_Dan_Van_Hanh_He_Thong_Va_Tro_Ly_AI",
        "title": "HƯỚNG DẪN KỸ THUẬT VẬN HÀNH HỆ THỐNG VÀ TRỢ LÝ ẢO AI CHATBOX",
        "so_hieu": "08/2026/HD-CNTT",
        "sections": [
            {
                "heading": "Chương I: KIẾN TRÚC TRỢ LÝ AI VÀ CHỈ MỤC CHROMA-DB",
                "content": (
                    "Điều 1. Kiến trúc mô hình ngôn ngữ và Vector Store\n"
                    "Trợ lý ảo AI TLBM vận hành trên nền tảng kỹ thuật RAG tiên tiến:\n"
                    "- LLM Model mặc định: Google Gemini 3.8 Flash (tích hợp cơ chế tự động fallback khi nghẽn mạng).\n"
                    "- Vector Database: ChromaDB lưu trữ cục bộ tại đường dẫn `./data/chroma_db` bảo mật tuyệt đối không gửi dữ liệu ra cloud công cộng.\n"
                    "- Cơ chế phân đoạn tri thức Parent-Child Chunking:\n"
                    "  + Parent Chunk (Sơ cấp): Kích thước 1.500 ký tự (Overlap 300 ký tự) cung cấp ngữ cảnh đầy đủ cho câu trả lời.\n"
                    "  + Child Chunk (Thứ cấp): Kích thước 300 ký tự (Overlap 50 ký tự) dùng để tạo vector embedding tra cứu siêu tốc.\n\n"
                    "Điều 2. Nguyên tắc Không Ảo giác (Zero-Hallucination)\n"
                    "AI chỉ được phản hồi dựa trên tài liệu hợp lệ trong kho tri thức. "
                    "Nếu không có dữ liệu nguồn, AI bắt buộc phải trả lời câu chuẩn: "
                    "'Xin lỗi, thông tin bạn hỏi không có trong tài liệu được cung cấp nên tôi không thể trả lời.'"
                ),
            },
            {
                "heading": "Chương II: LỊCH SAO LƯU DỮ LIỆU VÀ QUẢN TRỊ AN TOÀN THÔNG TIN",
                "content": (
                    "Điều 3. Chính sách sao lưu tự động (Backup Policy)\n"
                    "- Sao lưu hàng ngày: Toàn bộ CSDL PostgreSQL và thư mục `./data/chroma_db` được tự động sao lưu lúc 01:00 sáng hàng ngày.\n"
                    "- Sao lưu tuần: Đóng gói toàn bộ mã nguồn và dữ liệu vào lúc 02:00 sáng Chủ Nhật hàng tuần, lưu trữ tại máy chủ phụ ngoại vi.\n"
                    "- Thời hạn lưu trữ bản sao lưu: Tối thiểu 90 ngày trước khi xoay vòng dữ liệu.\n\n"
                    "Điều 4. Quyền quản trị phân hệ AI\n"
                    "Chỉ tài khoản Quản trị viên tối cao (is_superuser=True) mới có quyền truy cập trang Cài đặt AI, "
                    "xóa trắng chỉ mục Vector, hoặc thay đổi khóa API Key và Model Provider trong hệ thống."
                ),
            },
        ],
    },
]


def create_docx(doc_data: dict, out_file: Path):
    doc = docx.Document()

    # Căn lề chuẩn văn bản hành chính: Trên 2cm, Dưới 2cm, Trái 3cm, Phải 2cm
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.79)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(1.18)
        s.right_margin = Inches(0.79)

    # Quốc hiệu, Tiêu ngữ
    header_para = doc.add_paragraph()
    header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_qh = header_para.add_run("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n")
    run_qh.bold = True
    run_qh.font.size = Pt(13)
    run_line = header_para.add_run("-----------------------")
    run_line.bold = True

    # Tiêu đề tài liệu
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run(f"\n{doc_data['title']}\n")
    title_run.bold = True
    title_run.font.size = Pt(15)
    title_run.font.color.rgb = RGBColor(16, 44, 87)

    # Số hiệu
    sh_para = doc.add_paragraph()
    sh_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sh_run = sh_para.add_run(f"Số hiệu: {doc_data['so_hieu']} | Ngày ban hành: 10/10/2026\n")
    sh_run.italic = True
    sh_run.font.size = Pt(11)

    # Nội dung các chương / điều
    for sec in doc_data["sections"]:
        h_para = doc.add_paragraph()
        h_run = h_para.add_run(f"\n{sec['heading']}")
        h_run.bold = True
        h_run.font.size = Pt(13)
        h_run.font.color.rgb = RGBColor(26, 82, 118)

        c_para = doc.add_paragraph()
        c_para.paragraph_format.line_spacing = 1.25
        c_run = c_para.add_run(sec["content"])
        c_run.font.size = Pt(12)

    doc.save(str(out_file))


def create_txt(doc_data: dict, out_file: Path):
    lines = [
        "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
        "Độc lập - Tự do - Hạnh phúc",
        "-----------------------",
        "",
        doc_data["title"],
        f"Số hiệu: {doc_data['so_hieu']} | Ngày ban hành: 10/10/2026",
        "",
    ]
    for sec in doc_data["sections"]:
        lines.append(sec["heading"])
        lines.append(sec["content"])
        lines.append("")

    out_file.write_text("\n".join(lines), encoding="utf-8")


def main():
    print("=" * 70)
    print("📂  KHỞI TẠO BỘ TÀI LIỆU KIỂM THỬ TRỢ LÝ AI Ở NHIỀU THƯ MỤC KHÁC NHAU")
    print("=" * 70)

    TEST_DOCS_ROOT.mkdir(parents=True, exist_ok=True)
    created_files = []

    for doc_item in DOCUMENTS_DATA:
        folder_dir = TEST_DOCS_ROOT / doc_item["folder"]
        folder_dir.mkdir(parents=True, exist_ok=True)

        docx_path = folder_dir / f"{doc_item['filename_base']}.docx"
        txt_path = folder_dir / f"{doc_item['filename_base']}.txt"

        create_docx(doc_item, docx_path)
        create_txt(doc_item, txt_path)

        created_files.append((folder_dir.name, docx_path, doc_item["category"]))
        created_files.append((folder_dir.name, txt_path, doc_item["category"]))

        print(f"\n📁 Thư mục: {folder_dir.name}/ (Chủ đề: {doc_item['category']})")
        print(f"   📄 Tệp Word DOCX: {docx_path.name} ({docx_path.stat().st_size:,} bytes)")
        print(f"   📄 Tệp Text TXT:  {txt_path.name} ({txt_path.stat().st_size:,} bytes)")

    # Đồng thời sao chép các tệp vào data/docs để nạp sẵn vào kho AI nếu cần
    data_docs_dir = ROOT_DIR / "data" / "docs"
    data_docs_dir.mkdir(parents=True, exist_ok=True)

    print("\n" + "=" * 70)
    print("📥  ĐỒNG BỘ TÀI LIỆU VÀO KHO DỮ LIỆU CỦA AI CHATBOX (data/docs)...")
    print("=" * 70)

    # Đọc registry hiện tại
    registry_file = ROOT_DIR / "data" / "doc_registry.json"
    registry_list = []
    if registry_file.exists():
        try:
            registry_list = json.loads(registry_file.read_text(encoding="utf-8"))
        except Exception:
            registry_list = []

    existing_filenames = {e["filename"] for e in registry_list}

    for folder_name, fpath, category in created_files:
        # Chỉ copy tệp DOCX vào kho data/docs để tránh trùng lặp
        if fpath.suffix == ".docx":
            dest_file = data_docs_dir / fpath.name
            shutil.copy2(fpath, dest_file)
            print(f"   ✅ Đã copy vào data/docs/: {fpath.name}")

            if fpath.name not in existing_filenames:
                registry_list.append({
                    "filename": fpath.name,
                    "category": category,
                    "size": dest_file.stat().st_size,
                    "upload_time": datetime.now(timezone.utc).isoformat(),
                })
                existing_filenames.add(fpath.name)

    registry_file.write_text(json.dumps(registry_list, ensure_ascii=False, indent=2), encoding="utf-8")
    print("   ✅ Đã cập nhật data/doc_registry.json thành công.")

    print("\n" + "=" * 70)
    print(f"🎉 TỔNG KẾT: Đã tạo thành công {len(DOCUMENTS_DATA)} bộ tài liệu tại:")
    print(f"   👉 {TEST_DOCS_ROOT}")
    print("=" * 70)


if __name__ == "__main__":
    main()
